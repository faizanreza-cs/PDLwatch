"""PDLs published as spreadsheets (Excel or CSV), read as pages of positioned rows.

Some states publish their list as a workbook as well as, or instead of, a PDF (Washington posts
its Apple Health PDL every week as Excel). Each sheet becomes pages of lines: one line per row,
one phrase per filled cell, each column at its own position. The table parser then reads it
exactly as it reads a ruled PDF table: the same header vocabulary, legends, matching, scoring
and review, and every row cites its sheet ("page") and row position.

A workbook the header vocabulary can't read as it is takes a source profile (``profile:`` in the
sources file): the columns to keep, renamed to header words the parser knows, and the rows to
keep. Washington's lists every package (NDC) in 23 coded columns; its profile keeps the pharmacy
benefit's status and PA of the drugs on the PDL, and the packages of a product collapse into one
row. ``columns()`` lists each column's values, to write a profile from.
"""

from __future__ import annotations

import csv
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from pdlwatch.pdf import Line, Page, Phrase, Word

SUFFIXES = {".xlsx", ".xlsm", ".csv"}
ROWS_PER_PAGE = 60
COLUMN_WIDTH = 120.0
ROW_HEIGHT = 20.0  # far enough apart that no two rows read as one wrapped header


Profile = dict[str, dict[str, str]]
PROFILE_KEYS = ("columns", "rows")


def is_sheet(path: str | Path) -> bool:
    return Path(path).suffix.lower() in SUFFIXES


def _name(text: Any) -> str:
    return " ".join(str(text).upper().split())


def check_profile(profile: Any) -> list[str]:
    """What is wrong with a source profile, [] when nothing: its shape, and new column names the
    table parser wouldn't read as a header."""
    from pdlwatch.extract import layout

    if not isinstance(profile, dict) or not profile.get("columns"):
        return ["a profile needs columns: {spreadsheet header: header the parser reads}"]
    errors = [f"unknown profile key {k!r}" for k in profile if k not in PROFILE_KEYS]
    for key in PROFILE_KEYS:
        given = profile.get(key, {})
        if not isinstance(given, dict) or not all(
            isinstance(k, str) and isinstance(v, str) for k, v in given.items()
        ):
            errors.append(f"profile {key} must map names to text (quote Y and N)")
        elif key == "columns":
            errors += [f"{v!r} is not a header the parser reads" for v in given.values()
                       if layout.role(v) is None]  # fmt: skip
    return errors


def _profiled(rows: list[list[str]], profile: Profile) -> list[list[str]]:
    """A sheet as its source profile reads it. The rows above the header stay as they are (the
    list's title and date); below it come the profile's columns under their new names, only the
    rows its filters keep, each distinct row once. A sheet without the profile's header (a notes
    sheet) is returned as it is."""
    cols = {_name(k): v for k, v in profile["columns"].items()}
    where = {_name(k): _name(v) for k, v in profile.get("rows", {}).items()}
    at = next((i for i, row in enumerate(rows) if {*cols, *where} <= {_name(c) for c in row}),
              None)  # fmt: skip
    if at is None:
        return rows
    names = [_name(c) for c in rows[at]]
    pick = [names.index(c) for c in cols]
    tests = [(names.index(c), v) for c, v in where.items()]
    out = [*rows[:at], list(cols.values())]
    seen: set[tuple[str, ...]] = set()
    for row in rows[at + 1 :]:
        cells = row + [""] * (len(names) - len(row))
        if any(_name(cells[j]) != v for j, v in tests):
            continue
        kept = tuple(cells[j] for j in pick)
        if any(kept) and kept not in seen:  # every package of a product line: one row
            seen.add(kept)
            out.append(list(kept))
    return out


def columns(path: str | Path, top: int = 5) -> list[str]:
    """Each sheet's columns: the header, how many different values the column holds and the most
    common ones, to write a source profile from. The header is the opening row with the most
    filled cells."""
    out = []
    for sheet, rows in _tables(Path(path)):
        at = max(range(min(len(rows), 30)), key=lambda i: (sum(map(bool, rows[i])), -i))
        body = rows[at + 1 :]
        out.append(f"sheet {sheet}: {len(body):,} rows under the header on row {at + 1}")
        for j, head in enumerate(rows[at]):
            values = Counter(r[j] if j < len(r) else "" for r in body)
            common = ", ".join(f"{(v or '(blank)')[:40]} {n:,}" for v, n in values.most_common(top))
            out.append(f"  {head or f'(column {j + 1})'}: {len(values):,} values; {common}")
    return out


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return " ".join(str(value).split())


def _tables(path: Path) -> list[tuple[str, list[list[str]]]]:
    """(sheet name, rows of cell text) for every sheet with content."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig", errors="replace") as f:
            return [(path.stem, [[_text(v) for v in row] for row in csv.reader(f)])]
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        out = []
        for ws in wb.worksheets:
            rows = [[_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
            if any(any(row) for row in rows):
                out.append((ws.title, rows))
        return out
    finally:
        wb.close()


def _line(cells: list[str], top: float) -> Line | None:
    phrases = []
    for col, text in enumerate(cells):
        if not text:
            continue
        x = 20.0 + col * COLUMN_WIDTH
        # a cell's words stay inside its column however long the text: the parser assigns each
        # word to the column its left edge falls in
        words = [Word(w, x + i * 0.01, x + i * 0.01 + 0.005, top, top + 9.0)
                 for i, w in enumerate(text.split())]  # fmt: skip
        phrases.append(Phrase(words))
    return Line(phrases) if phrases else None


def load(path: str | Path, profile: Profile | None = None) -> list[Page]:
    """Every sheet as pages of up to ROWS_PER_PAGE rows, numbered on from one sheet to the
    next, read through the source's ``profile`` when it has one. The rows sit inside the page
    body, clear of the title and footer bands."""
    pages: list[Page] = []
    for _sheet, rows in _tables(Path(path)):
        if profile:
            rows = _profiled(rows, profile)
        width = 40.0 + COLUMN_WIDTH * max((len(r) for r in rows), default=1)
        for start in range(0, len(rows), ROWS_PER_PAGE):
            chunk = rows[start : start + ROWS_PER_PAGE]
            height = len(chunk) * ROW_HEIGHT / 0.8
            lines = [ln for i, cells in enumerate(chunk)
                     if (ln := _line(cells, 0.1 * height + i * ROW_HEIGHT))]  # fmt: skip
            if lines:
                pages.append(Page(len(pages) + 1, width, height, lines, False))
    return pages
